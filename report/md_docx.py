"""
md_docx.py - a small converter from the markdown used for the stage chapters and report/src to python-docx.

Handles: headings, paragraphs, bullet and numbered lists, pipe tables, images with captions, "Table N. caption" lines,
bold and italic text, and <!--docs-only--> blocks (dropped). Code spans are written as plain text.
"""
import re
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

NAVY = RGBColor(0x1F, 0x4E, 0x79)
FONT = "Times New Roman"
FIG_ROOT = Path(__file__).resolve().parent.parent / "hvac_ml" / "outputs" / "figures"
USABLE_W_IN = 6.3
MAX_IMG_H_IN = 7.6


# low-level helpers
def add_bookmark(paragraph, name, bid):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    paragraph._p.insert(0 if len(paragraph._p) == 0 or paragraph._p[0].tag != qn("w:pPr") else 1, start)
    paragraph._p.append(end)


def add_field(paragraph, instr, placeholder="0", bold=False, size=None):
    """Insert a complex field (PAGEREF, TOC, PAGE ...) with a placeholder result."""
    def run_with(child):
        r = paragraph.add_run()
        r.font.name = FONT
        if bold:
            r.bold = True
        if size:
            r.font.size = Pt(size)
        r._r.append(child)
        return r
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    run_with(b)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run_with(it)
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    run_with(s)
    r = paragraph.add_run(placeholder)
    r.font.name = FONT
    if bold:
        r.bold = True
    if size:
        r.font.size = Pt(size)
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    run_with(e)


def set_cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_cell_margins(table, top=40, bottom=40, left=70, right=70):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tblPr.append(m)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trPr.append(h)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    c = OxmlElement("w:cantSplit")
    c.set(qn("w:val"), "true")
    trPr.append(c)


def style_run(run, size=None, bold=None, italic=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")


def add_inline(paragraph, text, size=None, bold=None, italic=None, color=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = paragraph.add_run(part[2:-2])
            style_run(r, size, True, italic, color)
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            style_run(r, size, bold, italic, color)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = paragraph.add_run(part[1:-1])
            style_run(r, size, bold, True, color)
        else:
            r = paragraph.add_run(part)
            style_run(r, size, bold, italic, color)


def setup_styles(doc):
    st = doc.styles
    n = st["Normal"]
    n.font.name = FONT
    n.font.size = Pt(11)
    n.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.15
    for name, size, color, before, after in (("Heading 1", 18, NAVY, 0, 12), ("Heading 2", 13.5, NAVY, 14, 5), ("Heading 3", 11.5, RGBColor(0x2E, 0x86, 0xAB), 10, 3)):
        h = st[name]
        h.font.name = FONT
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.italic = False
        h.font.color.rgb = color
        h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        h.element.rPr.rFonts.set(qn("w:ascii"), FONT)
        h.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if h.element.rPr.rFonts.get(qn(att)) is not None:
                del h.element.rPr.rFonts.attrib[qn(att)]
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True
    st["Heading 1"].paragraph_format.page_break_before = True
    c = st["Caption"]
    c.font.name = FONT
    c.font.size = Pt(10)
    c.font.italic = True
    c.font.bold = False
    c.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    c.paragraph_format.space_after = Pt(10)


# the converter
class Converter:
    def __init__(self, doc):
        self.doc = doc
        self.bm_id = 100
        self.figures = []        # (label, caption, bookmark)
        self.chapters = {}       # chapter number -> (start bookmark, end bookmark)

    def next_bm(self):
        self.bm_id += 1
        return self.bm_id

    # ---- block level
    def convert(self, md_text, chapter, owner_line=None):
        text = re.sub(r"<!--docs-only-->.*?<!--/docs-only-->", "", md_text, flags=re.S)
        lines = text.splitlines()
        i, fig_n, last_par = 0, 0, None
        first_heading = True
        pending_table_caption = None
        while i < len(lines):
            line = lines[i].rstrip()
            if not line.strip():
                i += 1
                continue
            m = re.match(r"^(#{1,3}) (.+)$", line)
            if m:
                level, title = len(m.group(1)), m.group(2).strip()
                title = re.sub(r"^Stage (\d+)\.", r"Chapter \1.", title)
                p = self.doc.add_heading(title, level=level)
                if level == 1 and chapter is not None:
                    add_bookmark(p, f"ch{chapter}_start", self.next_bm())
                    self.chapters[chapter] = [f"ch{chapter}_start", None]
                last_par = p
                i += 1
                continue
            if line.startswith("Owner: "):
                who = re.match(r"Owner: (.+?)\.", line)
                p = self.doc.add_paragraph()
                add_inline(p, f"*Lifecycle stage {chapter} of 10. Written by {who.group(1)}.*" if who else line)
                last_par = p
                i += 1
                continue
            mi = re.match(r"^!\[(.+?)\]\((.+?)\)$", line)
            if mi:
                fig_n += 1
                last_par = self.add_figure(mi.group(1), mi.group(2), chapter, fig_n)
                i += 1
                continue
            mt = re.match(r"^Table (\d+)\. (.+)$", line)
            if mt and i + 2 < len(lines) and lines[i + 2].lstrip().startswith("|"):
                pending_table_caption = f"Table {chapter}.{mt.group(1)}. {mt.group(2)}"
                i += 1
                continue
            if line.lstrip().startswith("|"):
                rows = []
                while i < len(lines) and lines[i].lstrip().startswith("|"):
                    rows.append(lines[i].strip())
                    i += 1
                last_par = self.add_table(rows, pending_table_caption)
                pending_table_caption = None
                continue
            ml = re.match(r"^\s*([*-]|\d+\.) (.+)$", line)
            if ml:
                items = []
                while i < len(lines):
                    mm = re.match(r"^\s*([*-]|\d+\.) (.+)$", lines[i].rstrip())
                    if not mm:
                        break
                    marker, body = mm.group(1), mm.group(2)
                    j = i + 1
                    while j < len(lines) and lines[j].startswith("   ") and lines[j].strip():   # continuation line
                        body += " " + lines[j].strip()
                        j += 1
                    items.append((marker, body))
                    i = j
                for k, (marker, body) in enumerate(items):
                    p = self.doc.add_paragraph()
                    pf = p.paragraph_format
                    pf.left_indent, pf.first_line_indent, pf.space_after = Cm(0.9), Cm(-0.6), Pt(3)
                    pf.tab_stops.add_tab_stop(Cm(0.9))
                    label = "•" if marker in "*-" else marker
                    r = p.add_run(label + "\t")
                    style_run(r)
                    add_inline(p, self.fix_refs(body, chapter))
                    last_par = p
                self.doc.paragraphs[-1].paragraph_format.space_after = Pt(8)
                continue
            # paragraph (merge following non-blank plain lines)
            body = line
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||!\[|\s*([*-]|\d+\.) )", lines[i]) and not lines[i].startswith("Table "):
                body += " " + lines[i].strip()
                i += 1
            p = self.doc.add_paragraph()
            add_inline(p, self.fix_refs(body, chapter))
            last_par = p
        if chapter is not None and chapter in self.chapters and last_par is not None:
            add_bookmark(last_par, f"ch{chapter}_end", self.next_bm())
            self.chapters[chapter][1] = f"ch{chapter}_end"

    def fix_refs(self, text, chapter):
        if chapter is None:
            return text
        text = re.sub(r"\bFigure (\d+)\b", lambda m: f"Figure {chapter}.{m.group(1)}", text)
        text = re.sub(r"\bTable (\d+)\b", lambda m: f"Table {chapter}.{m.group(1)}", text)
        return text

    # ---- figures
    def add_figure(self, caption, rel_path, chapter, n):
        path = self.find_image(rel_path)
        with Image.open(path) as im:
            w_px, h_px = im.size
        w_in = USABLE_W_IN
        h_in = w_in * h_px / w_px
        if h_in > MAX_IMG_H_IN:
            w_in, h_in = MAX_IMG_H_IN * w_px / h_px, MAX_IMG_H_IN
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(path), width=Inches(w_in))
        label = f"Figure {chapter}.{n}" if chapter is not None else f"Figure {n}"
        text = re.sub(r"^Figure \d+\.\s*", "", caption)
        c = self.doc.add_paragraph(style="Caption")
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.keep_together = True
        r = c.add_run(f"{label}. ")
        style_run(r, 10, True, True)
        add_inline(c, text, size=10, italic=True)
        bm = f"fig_{chapter}_{n}"
        add_bookmark(c, bm, self.next_bm())
        self.figures.append((label, text, bm))
        return c

    def find_image(self, rel):
        name = Path(rel).name
        for d in (FIG_ROOT / "report", FIG_ROOT):
            if (d / name).exists():
                return d / name
        raise FileNotFoundError(rel)

    # ---- tables
    def add_table(self, rows, caption):
        cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        header, body = cells[0], [r for r in cells[2:]]
        ncol = len(header)
        body = [(r + [""] * ncol)[:ncol] for r in body]
        if caption:
            cp = self.doc.add_paragraph(style="Caption")
            cp.paragraph_format.keep_with_next = True
            cp.paragraph_format.space_before = Pt(8)
            cp.paragraph_format.space_after = Pt(3)
            cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            lab, rest = caption.split(". ", 1)
            r = cp.add_run(lab + ". ")
            style_run(r, 10, True, True)
            add_inline(cp, rest, size=10, italic=True)
        font = 10 if ncol <= 3 else 9.5 if ncol <= 4 else 9 if ncol <= 6 else 8 if ncol <= 8 else 7.5
        charw = 0.0075 * font                      # inches per character, Times New Roman
        pad = 0.16
        def tok(t):
            return re.sub(r"[*`]", "", t).replace(" ± ", " ± ")
        need, pref = [], []
        for j in range(ncol):
            texts = [tok(row[j]) for row in [header] + body]
            need.append(max(max((len(w) for w in t.replace(" ", "~").split()), default=1) for t in texts) * charw + pad)
            pref.append(min(max(len(t) for t in texts) * charw + pad, 2.6))
        pref = [max(p_, n_) for p_, n_ in zip(pref, need)]
        if sum(pref) <= USABLE_W_IN:
            widths = [p_ * USABLE_W_IN / sum(pref) for p_ in pref]
        elif sum(need) >= USABLE_W_IN:
            widths = [n_ * USABLE_W_IN / sum(need) for n_ in need]
        else:
            spare = USABLE_W_IN - sum(need)
            extra = [p_ - n_ for p_, n_ in zip(pref, need)]
            widths = [n_ + spare * e_ / sum(extra) for n_, e_ in zip(need, extra)]
        body = [[c.replace(" ± ", " ± ") for c in r] for r in body]
        header = [c.replace(" ± ", " ± ") for c in header]
        t = self.doc.add_table(rows=1 + len(body), cols=ncol)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        set_cell_margins(t)
        for ri, row in enumerate([header] + body):
            tr = t.rows[ri]
            no_split(tr)
            if ri == 0:
                repeat_header(tr)
            for cj, txt in enumerate(row):
                cell = tr.cells[cj]
                cell.width = Inches(widths[cj])
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                if ri == 0:
                    set_cell_shading(cell, "1F4E79")
                    add_inline(p, txt, size=font, bold=True, color=RGBColor(255, 255, 255))
                else:
                    if ri % 2 == 0:
                        set_cell_shading(cell, "F2F5F8")
                    add_inline(p, txt, size=font)
                    if cj > 0 and re.fullmatch(r"[-+−]?[\d.,]+([  ]?[%±][  \d.,]*)?", txt.strip()):
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if len(body) <= 9:
            for tr in t.rows[:-1]:
                for c in tr.cells:
                    for pp in c.paragraphs:
                        pp.paragraph_format.keep_with_next = True
        after = self.doc.add_paragraph()
        after.paragraph_format.space_after = Pt(4)
        after.paragraph_format.line_spacing = 0.6
        return after
