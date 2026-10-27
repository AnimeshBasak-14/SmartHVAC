"""
build_report.py - builds report/HVAC_Service_Planner_Report.docx from the stage chapters in team_docs/docs/lifecycle/*.md (kept outside this repository, next to it) and report/src/*.md.

    python report/build_report.py
Then run report/finish_docx.ps1 (Word) to update the table of contents and page numbers and export a PDF.
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from md_docx import FONT, NAVY, Converter, add_bookmark, add_field, add_inline, repeat_header, set_cell_margins, set_cell_shading, setup_styles, style_run  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
LIFE = ROOT.parent / "team_docs" / "docs" / "lifecycle"
SRC = ROOT / "report" / "src"
OUT = ROOT / "report" / "HVAC_Service_Planner_Report.docx"

MEMBERS = [
    ("Kriti Vajpayee", "25931008", [1, 2], "Framed the problem and the product (stage 1). Searched for, downloaded and documented the three datasets and their sources (stage 2)."),
    ("Dhruv Aggarwal", "25911019", [3, 4], "Cleaned the data, built the churn table, designed the train and test splits and the leakage tests (stage 3). Explored the data and drew the plots that guided the features (stage 4)."),
    ("Ankit Kumar", "26911002", [5, 6], "Designed and tested the engineered features (stage 5). Chose the candidate models and the selection rule, and ran the model contest (stage 6)."),
    ("Animesh Basak", "26911001", [7, 8], "Trained all models under one protocol and measured learning curves and sensitivity (stage 7). Ran the held-out evaluation, error analysis and calibration checks (stage 8)."),
    ("Sourav Kumar Boro", "26535026", [9, 10], "Built and tested the dashboard and the heatwave safeguard (stage 9). Wrote the monitoring checks, alert rules and maintenance plan (stage 10)."),
]
CHAPTER_FILES = [
    (1, "01_Problem_Definition.md"), (2, "02_Data_Collection.md"), (3, "03_Data_Cleaning_and_Preprocessing.md"), (4, "04_Exploratory_Data_Analysis.md"),
    (5, "05_Feature_Engineering.md"), (6, "06_Model_Selection.md"), (7, "07_Model_Training.md"), (8, "08_Model_Evaluation_and_Tuning.md"),
    (9, "09_Model_Deployment.md"), (10, "10_Monitoring_and_Maintenance.md"),
]
TAIL_FILES = [(11, "11_Results_and_Discussion.md"), (12, "12_Conclusion_and_Recommendation.md"), (13, "13_Limitations_and_Future_Work.md"), (None, "14_References.md")]


def centered(doc, text, size, bold=False, color=None, before=0, after=6, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
    r = p.add_run(text)
    style_run(r, size, bold, italic, color)
    return p


def cover(doc):
    centered(doc, "Indian Institute of Technology Roorkee", 20, True, NAVY, before=30, after=2)
    centered(doc, "Department of Computer Science and Engineering", 14, False, None, after=40)
    centered(doc, "HVAC Service Planner", 32, True, NAVY, after=6)
    centered(doc, "Using Machine Learning for HVAC System Providers such as Aire Serv", 16, False, None, after=6, italic=True)
    centered(doc, "Demand forecasting, failure prediction and customer retention", 13, False, RGBColor(0x55, 0x55, 0x55), after=36)
    centered(doc, "Project Report", 16, True, None, after=4)
    centered(doc, "Course: Machine Learning (CSC-505)", 13, False, None, after=2)
    centered(doc, "Group 1", 13, True, None, after=22)
    t = doc.add_table(rows=1 + len(MEMBERS), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    set_cell_margins(t, 70, 70, 120, 120)
    for j, h in enumerate(("Group member", "Enrollment number")):
        c = t.rows[0].cells[j]
        c.width = Inches(2.8 if j == 0 else 2.0)
        set_cell_shading(c, "1F4E79")
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        add_inline(c.paragraphs[0], h, 12, True, color=RGBColor(255, 255, 255))
    for i, (name, enr, _, _) in enumerate(MEMBERS, 1):
        for j, v in enumerate((name, enr)):
            c = t.rows[i].cells[j]
            c.width = Inches(2.8 if j == 0 else 2.0)
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            add_inline(c.paragraphs[0], v, 12)
            if j == 1:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    centered(doc, "Under the guidance of", 12, False, None, before=28, after=2)
    centered(doc, "Dr. Durga Toshniwal", 14, True, None, after=0)
    centered(doc, "Professor, Department of Computer Science and Engineering, IIT Roorkee", 12, False, None, after=30)
    centered(doc, "October 2026", 12, False, None, after=0)


def contributions(doc):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    r = p.add_run("Contribution of the group members")
    style_run(r, 18, True, False, NAVY)
    add_inline(doc.add_paragraph(), "Each member wrote the report chapters for two lifecycle stages and built the part of the project that those stages describe. "
               "The page numbers are the first and last page of each member's chapters.")
    t = doc.add_table(rows=1 + len(MEMBERS), cols=5)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_cell_margins(t, 60, 60, 80, 80)
    widths = [1.0, 0.8, 1.45, 2.3, 0.75]
    heads = ["Member", "Enrollment", "Lifecycle stages", "Contribution", "Pages"]
    stage_names = {1: "Problem definition", 2: "Data collection", 3: "Data cleaning and preprocessing", 4: "Exploratory data analysis", 5: "Feature engineering",
                   6: "Model selection", 7: "Model training", 8: "Model evaluation and tuning", 9: "Model deployment", 10: "Monitoring and maintenance"}
    for j, h in enumerate(heads):
        c = t.rows[0].cells[j]
        c.width = Inches(widths[j])
        set_cell_shading(c, "1F4E79")
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        add_inline(c.paragraphs[0], h, 10, True, color=RGBColor(255, 255, 255))
    repeat_header(t.rows[0])
    for i, (name, enr, stages, what) in enumerate(MEMBERS, 1):
        row = t.rows[i]
        for j in range(5):
            row.cells[j].width = Inches(widths[j])
            row.cells[j].paragraphs[0].paragraph_format.space_after = Pt(0)
            row.cells[j].paragraphs[0].paragraph_format.line_spacing = 1.0
            if i % 2 == 0:
                set_cell_shading(row.cells[j], "F2F5F8")
        add_inline(row.cells[0].paragraphs[0], name, 10, True)
        add_inline(row.cells[1].paragraphs[0], enr, 10)
        add_inline(row.cells[2].paragraphs[0], "; ".join(f"{s}. {stage_names[s]}" for s in stages), 10)
        add_inline(row.cells[3].paragraphs[0], what, 10)
        pp = row.cells[4].paragraphs[0]
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(pp, f"PAGEREF ch{stages[0]}_start \\h", "0", size=10)
        r = pp.add_run(" to ")
        style_run(r, 10)
        add_field(pp, f"PAGEREF ch{stages[1]}_end \\h", "0", size=10)
    doc.add_paragraph()
    add_inline(doc.add_paragraph(), "**Shared work.** All five members took part in choosing the topic and the product, in reviewing each other's chapters, and in preparing the presentation. "
               "The report chapters 11 to 13 (results, conclusion and recommendation, limitations and future work) were written jointly.")
    add_inline(doc.add_paragraph(), "**Where to find the other parts.** Results and discussion, conclusion with the recommendation, limitations and future work, and the references follow the ten lifecycle chapters.")
    t2 = doc.add_table(rows=5, cols=3)
    t2.style = "Table Grid"
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t2.autofit = False
    set_cell_margins(t2, 50, 50, 80, 80)
    rows = [("Part of the report", "Where", "Page"), ("Abstract", "abstract", None), ("Results and discussion", "ch11_start", None), ("Conclusion and recommendation", "ch12_start", None), ("References and data sources", "refs_start", None)]
    for i, (a, b, _) in enumerate(rows):
        for j, w in enumerate((3.2, 1.5, 0.8)):
            t2.rows[i].cells[j].width = Inches(w)
            t2.rows[i].cells[j].paragraphs[0].paragraph_format.space_after = Pt(0)
        if i == 0:
            for j, v in enumerate(rows[0]):
                set_cell_shading(t2.rows[0].cells[j], "1F4E79")
                add_inline(t2.rows[0].cells[j].paragraphs[0], v, 10, True, color=RGBColor(255, 255, 255))
            continue
        add_inline(t2.rows[i].cells[0].paragraphs[0], a, 10)
        add_inline(t2.rows[i].cells[1].paragraphs[0], "Chapter " + b[2:4] if b.startswith("ch") else "Front matter" if b == "abstract" else "End", 10)
        t2.rows[i].cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(t2.rows[i].cells[2].paragraphs[0], f"PAGEREF {b} \\h", "0", size=10)


def footer_and_header(section):
    section.different_first_page_header_footer = True
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hp.add_run("HVAC Service Planner  |  Group 1  |  CSC-505 Machine Learning, IIT Roorkee")
    style_run(r, 9, False, True, RGBColor(0x66, 0x66, 0x66))
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, "PAGE", "1", size=10)


def toc_page(doc):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    r = p.add_run("Table of contents")
    style_run(r, 18, True, False, NAVY)
    add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u', "Right-click and choose Update field to build the table of contents.")


def list_of_figures(doc, conv):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    r = p.add_run("List of figures")
    style_run(r, 18, True, False, NAVY)
    for label, text, bm in conv.figures:
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(1)
        q.paragraph_format.left_indent = Cm(1.9)
        q.paragraph_format.first_line_indent = Cm(-1.9)
        q.paragraph_format.tab_stops.add_tab_stop(Cm(1.9))
        q.paragraph_format.tab_stops.add_tab_stop(Cm(15.9), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r1 = q.add_run(label.replace("Figure ", "") + "\t")
        style_run(r1, 9.5, True)
        r2 = q.add_run(text[:105] + ("..." if len(text) > 105 else "") + "\t")
        style_run(r2, 9.5)
        add_field(q, f"PAGEREF {bm} \\h", "0", size=9.5)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2.2)
    setup_styles(doc)
    footer_and_header(sec)
    props = doc.core_properties
    props.title, props.author, props.subject = "HVAC Service Planner: project report", "Group 1, CSC-505", "Machine Learning (CSC-505), IIT Roorkee"

    cover(doc)
    contributions(doc)
    conv = Converter(doc)
    # abstract
    conv.convert((SRC / "00_abstract.md").read_text(encoding="utf-8"), None)
    ab = [p for p in doc.paragraphs if p.text == "Abstract"][0]
    add_bookmark(ab, "abstract", conv.next_bm())
    toc_page(doc)
    lof_anchor = doc.add_paragraph()      # list of figures is inserted here once every figure is known
    for ch, fn in CHAPTER_FILES:
        conv.convert((LIFE / fn).read_text(encoding="utf-8"), ch)
    for ch, fn in TAIL_FILES:
        before = len(doc.paragraphs)
        conv.convert((SRC / fn).read_text(encoding="utf-8"), ch)
        if ch is None:
            add_bookmark(doc.paragraphs[before], "refs_start", conv.next_bm())
    # build the list of figures at the end, then move it to the anchor
    start_idx = len(doc.element.body)
    list_of_figures(doc, conv)
    body = doc.element.body
    moved = list(body)[start_idx - 1:-1]       # elements added by list_of_figures (sectPr stays last)
    anchor = lof_anchor._p
    for el in reversed(moved):
        anchor.addnext(el)
    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print("saved", OUT, "figures:", len(conv.figures))


if __name__ == "__main__":
    main()
