import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def build_presentation():
    template_path = "docs/New folder/IITR_template_pptx(19062016)/IITR_template.pptx"
    out_path = "docs/New folder/SmartHVAC_Full_Project_PPT.pptx"

    prs = pptx.Presentation(template_path)

    # Helper styling function
    def set_font(run, text, name="Times New Roman", size_pt=14, bold=False, color=(30, 30, 30)):
        run.text = text
        run.font.name = name
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(*color)

    # Remove template placeholder slides after first
    slide_ids = [slide.slide_id for slide in prs.slides]
    for s_id in slide_ids[1:]:
        rId = prs.slides._sldIdLst.find(f'{{http://schemas.openxmlformats.org/presentationml/2006/main}}sldId[@id="{s_id}"]')
        if rId is not None:
            prs.part.drop_rel(rId.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'])
            prs.slides._sldIdLst.remove(rId)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Modify Slide 0)
    # -------------------------------------------------------------
    s0 = prs.slides[0]
    for shape in s0.shapes:
        if shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            if "Title" in shape.name or shape.top < Inches(3.0):
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                r = p.add_run()
                set_font(r, "Heterogeneous Graph Neural Network for Smart HVAC Energy & Appliance Management", "Times New Roman", 22, bold=True, color=(15, 34, 64))
                
                p2 = tf.add_paragraph()
                p2.alignment = PP_ALIGN.CENTER
                r2 = p2.add_run()
                set_font(r2, "Predictive Maintenance, Customer Churn, Energy Forecasting & Supply Chain Optimization", "Times New Roman", 13, bold=True, color=(31, 78, 121))
            else:
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                r = p.add_run()
                set_font(r, "Course: Machine Learning (CSL-505) / Data Mining (CSL-515)\nDepartment of Computer Science & Engineering, IIT Roorkee\n", "Times New Roman", 12, bold=True, color=(44, 62, 80))
                
                p_team = tf.add_paragraph()
                p_team.alignment = PP_ALIGN.CENTER
                r_team = p_team.add_run()
                set_font(r_team, "Project Team: Kriti | Dhruv | Ankit | Animesh Basak (26911001) | Saurav\nUnder Guidance of: Dr. Durga Toshniwal, Professor, CSE, IIT Roorkee", "Times New Roman", 11.5, bold=False, color=(50, 50, 50))

    # Helper function to add content slide
    def add_content_slide(title_text):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        if slide.shapes.title:
            slide.shapes.title.text = title_text
            for p in slide.shapes.title.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(20)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(15, 34, 64)
        return slide

    # -------------------------------------------------------------
    # SLIDE 2: Members' Contribution
    # -------------------------------------------------------------
    s2 = add_content_slide("Chapter 1: Group Members' Contribution")
    rows, cols = 6, 4
    left, top, width, height = Inches(0.8), Inches(1.8), Inches(8.4), Inches(4.5)
    tbl_shape = s2.shapes.add_table(rows, cols, left, top, width, height)
    tbl = tbl_shape.table
    tbl.columns[0].width = Inches(1.3)
    tbl.columns[1].width = Inches(1.8)
    tbl.columns[2].width = Inches(2.3)
    tbl.columns[3].width = Inches(3.0)

    headers = ["Member", "Enrollment / Role", "Lifecycle Phase", "Core Contributions & Scope"]
    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER if c < 2 else PP_ALIGN.LEFT
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11)
                r.font.bold = True
                r.font.color.rgb = RGBColor(15, 34, 64)

    data_team = [
        ("Kriti", "Team Member", "Problem Definition &\nData Collection", "Defined 4 operational pillars; collected 12-month climate feeds, device specs, and inventory data; built initial schema."),
        ("Dhruv", "Team Member", "Data Cleaning, Preprocessing\n& Exploratory Analysis (EDA)", "Thermodynamic sensor clipping, null handling, cyclic sin/cos encoding, acute class imbalance (<8%) analysis across regimes."),
        ("Ankit", "Team Member", "Feature Engineering &\nModel Selection", "Engineered rolling COP ratios, vibration stress, tenant RFM metrics; formulated RGCN / HeteroConv over tabular models."),
        ("Animesh Basak", "26911001", "Model Training, Evaluation\n& Fine-Tuning", "Tesla V100 GPU training (FP16 AMP, 28.8MB VRAM); Kendall uncertainty loss, Huber loss; UCI/Kaggle fine-tuning (-65.8% loss)."),
        ("Saurav", "Team Member", "Model Deployment &\nMonitoring / Maintenance", "Streamlit Cloud interactive dashboard (5 clean tabs); drift detection protocols (covariate & concept); automated retrain triggers.")
    ]

    for r_idx, row in enumerate(data_team, start=1):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx, c_idx)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER if c_idx < 2 else PP_ALIGN.LEFT
                for r in p.runs:
                    is_animesh = "Animesh" in row[0]
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(10)
                    r.font.bold = is_animesh
                    r.font.color.rgb = RGBColor(15, 34, 64) if is_animesh else RGBColor(40, 40, 40)

    # -------------------------------------------------------------
    # SLIDE 3: Introduction & Motivation
    # -------------------------------------------------------------
    s3 = add_content_slide("Chapter 2: Introduction & Motivation")
    tb3 = s3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(5.0))
    tf3 = tb3.text_frame
    tf3.word_wrap = True

    bullets3 = [
        ("The Energy Challenge: ", "Heating, Ventilation, and Air Conditioning (HVAC) accounts for over 40% of commercial and residential electrical power consumption in modern infrastructure."),
        ("The Fragmentation Problem: ", "Traditional facilities operate in strict functional silos: facility managers forecast thermal load via single-building physics (EnergyPlus), maintenance crews use calendar schedules, and inventory managers hold static spare buffers."),
        ("Cascading Climate Failures: ", "Extreme 45°C summer heatwaves or 90% humidity monsoons cause correlated multi-compressor failures, parts stockouts, and dissatisfied tenant subscription cancellations (sleeper churn)."),
        ("Why Graph Machine Learning (GML)? ", "Standard tabular ML (XGBoost) and time-series (LSTMs) treat buildings, appliances, residents, and spare parts as independent observations. Heterogeneous GNNs model the complete cyber-physical topology jointly.")
    ]
    for b_title, b_desc in bullets3:
        p = tf3.add_paragraph()
        p.space_after = Pt(12)
        r1 = p.add_run()
        set_font(r1, b_title, "Times New Roman", 13, bold=True, color=(15, 34, 64))
        r2 = p.add_run()
        set_font(r2, b_desc, "Times New Roman", 12, color=(40, 40, 40))

    # -------------------------------------------------------------
    # SLIDE 4: Problem Statement & The 4 Pillars (Kriti)
    # -------------------------------------------------------------
    s4 = add_content_slide("Chapter 3: Problem Statement - The 4 Pillars (Kriti)")
    tb4 = s4.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(5.0))
    tf4 = tb4.text_frame
    tf4.word_wrap = True

    pillars = [
        ("Pillar 1: Regional & Grid-Level Thermal Demand: ", "Forecast continuous heating and cooling servicing demand (kW/ton) across 12 distinct Indian climatic regimes (metro summer cooling peaks vs. rural winter heating surges)."),
        ("Pillar 2: Device Health & Predictive Maintenance: ", "Jointly predict continuous operating power draw (kW) and binary mechanical breakdown probability, catching capacitor failures and compressor burnouts before breakdown."),
        ("Pillar 3: Resident 'Sleeper' Churn Economics: ", "Classify off-season dormant residents (Recency > 90 days), calculating real Customer Lifetime Value (CLV) to trigger targeted maintenance tune-ups and prevent service churn."),
        ("Pillar 4: Spare Parts Supply Chain Optimization: ", "Forecast monthly SKU velocity and lead-time demand for mission-critical parts (capacitors, PCBs, valves) to eliminate stockouts while minimizing holding costs.")
    ]
    for p_title, p_desc in pillars:
        p = tf4.add_paragraph()
        p.space_after = Pt(10)
        r1 = p.add_run()
        set_font(r1, p_title, "Times New Roman", 13, bold=True, color=(31, 78, 121))
        r2 = p.add_run()
        set_font(r2, p_desc, "Times New Roman", 12, color=(40, 40, 40))

    # -------------------------------------------------------------
    # SLIDE 5: The Dataset & Heterogeneous Graph Schema (Kriti)
    # -------------------------------------------------------------
    s5 = add_content_slide("Chapter 4: The Dataset & Graph Schema (Kriti)")
    tb5 = s5.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    tf5 = tb5.text_frame
    tf5.word_wrap = True

    schema_points = [
        ("Unified Graph Topology: ", "4 distinct node types connected by 3 relational edge types across 12 monthly snapshots:"),
        ("• Location Nodes (40): ", "Ambient temp (°C), humidity (%), wind, solar radiation, urban heat index, building volume."),
        ("• Device Nodes (180): ", "Age, run hours, rolling COP, vibration RMS, motor current, refrigerant pressure, rated kW."),
        ("• Resident Nodes (600): ", "Income tier, occupancy count, thermostat setpoint, app engagement, billing credit score."),
        ("• Part Nodes (50): ", "Procurement cost, stock level, lead time, safety threshold, component MTBF."),
        ("Relational Edges: ", "(Location)-[:CONTAINS]->(Device), (Resident)-[:OWNS]->(Device), and (Device)-[:REQUIRES]->(Part).")
    ]
    for t_pt, d_pt in schema_points:
        p = tf5.add_paragraph()
        p.space_after = Pt(6)
        r1 = p.add_run()
        set_font(r1, t_pt, "Times New Roman", 11.5, bold=True, color=(15, 34, 64))
        r2 = p.add_run()
        set_font(r2, d_pt, "Times New Roman", 11, color=(50, 50, 50))

    img_schema = "visualization/02_heterogeneous_graph_schema.png"
    if os.path.exists(img_schema):
        s5.shapes.add_picture(img_schema, Inches(5.4), Inches(1.8), Inches(4.2), Inches(4.8))

    # -------------------------------------------------------------
    # SLIDE 6: Dataset Preprocessing & Cleaning (Dhruv)
    # -------------------------------------------------------------
    s6 = add_content_slide("Chapter 5: Dataset Preprocessing & Cleaning (Dhruv)")
    tb6 = s6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(5.0))
    tf6 = tb6.text_frame
    tf6.word_wrap = True

    prep_points = [
        ("Sensor Telemetry Sanitization: ", "Handled sensor dropouts and noise by applying rolling median filters and physics-based thermodynamic clipping (e.g., maximum electrical power capped at 3.5 kW nominal capacity)."),
        ("Multi-Modal Missing Imputation: ", "Imputed missing weather telemetry using spatial proximity averages from neighboring location nodes, and filled incomplete resident demographic profiles via median decile defaults."),
        ("Cyclical Temporal Encoding: ", "Transformed calendar months (1 to 12) into continuous 2D sinusoidal embeddings: theta_m = 2*pi*(m-1)/12, injecting sin(theta_m) and cos(theta_m) directly into node representations."),
        ("Graph Bipartite Sanitization: ", "Ensured 100% referential integrity across all bipartite edges, validating that every device connects to a valid resident owner, location envelope, and replacement parts catalog.")
    ]
    for b_title, b_desc in prep_points:
        p = tf6.add_paragraph()
        p.space_after = Pt(12)
        r1 = p.add_run()
        set_font(r1, b_title, "Times New Roman", 13, bold=True, color=(15, 34, 64))
        r2 = p.add_run()
        set_font(r2, b_desc, "Times New Roman", 12, color=(40, 40, 40))

    # -------------------------------------------------------------
    # SLIDE 7: Exploratory Data Analysis & Regimes (Dhruv)
    # -------------------------------------------------------------
    s7 = add_content_slide("Chapter 5: Exploratory Data Analysis & Climate Regimes (Dhruv)")
    tb7 = s7.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    tf7 = tb7.text_frame
    tf7.word_wrap = True

    eda_points = [
        ("Acute Class Imbalance (<8%): ", "Severe positive class rarity in failure events (2.07% to 7.8%) and off-season churn. Standard cross-entropy leads to trivial all-zero majority collapse."),
        ("Micro-Climate Disparities: ", "Metropolitan urban centers experience intense summer cooling loads peaking in May (624 units), whereas rural village hubs peak in winter heating demand (722 units in Dec)."),
        ("Monsoon PCB Moisture Surge: ", "High humidity (88-94%) during July-August causes electrical insulation breakdown and inverter shorts, shifting fault distributions dramatically from thermal overloads.")
    ]
    for b_title, b_desc in eda_points:
        p = tf7.add_paragraph()
        p.space_after = Pt(8)
        r1 = p.add_run()
        set_font(r1, b_title, "Times New Roman", 12, bold=True, color=(15, 34, 64))
        r2 = p.add_run()
        set_font(r2, b_desc, "Times New Roman", 11, color=(50, 50, 50))

    img_regimes = "visualization/09_seasonal_cooling_heating_regimes.png"
    if os.path.exists(img_regimes):
        s7.shapes.add_picture(img_regimes, Inches(5.4), Inches(2.0), Inches(4.2), Inches(4.4))

    # -------------------------------------------------------------
    # SLIDE 8: Feature Engineering & Selection (Ankit)
    # -------------------------------------------------------------
    s8 = add_content_slide("Chapter 6: Feature Engineering & Selection (Ankit)")
    tb8 = s8.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(5.0))
    tf8 = tb8.text_frame
    tf8.word_wrap = True

    feat_points = [
        ("Rolling Coefficient of Performance (COP): ", "Engineered dynamic thermal efficiency ratio COP = Q_thermal / P_elec. A decaying COP over consecutive months strongly signals refrigerant leakage or coil fouling."),
        ("Cumulative Vibration Stress Index: ", "Calculated RMS motor vibration velocity scaled by running hours. Captures mechanical bearing wear before catastrophic motor lockup."),
        ("Customer RFM & Sleeper Profiling: ", "Computed Recency (R), Frequency (F), and Monetary spend (M) from transaction histories. Customers with R > 90 days during shoulder seasons are tagged as 'Sleepers'."),
        ("Inventory Stockout Vulnerability: ", "Formulated dynamic reorder buffer ratio: Stockout_Risk = (Lead_Time * Demand_Velocity) / Current_Stock, highlighting SKUs facing urgent stockout risk.")
    ]
    for b_title, b_desc in feat_points:
        p = tf8.add_paragraph()
        p.space_after = Pt(11)
        r1 = p.add_run()
        set_font(r1, b_title, "Times New Roman", 12.5, bold=True, color=(15, 34, 64))
        r2 = p.add_run()
        set_font(r2, b_desc, "Times New Roman", 11.5, color=(40, 40, 40))

    # -------------------------------------------------------------
    # SLIDE 9: Model Selection & Relational Architecture (Ankit)
    # -------------------------------------------------------------
    s9 = add_content_slide("Chapter 7: Model Selection & Relational Architecture (Ankit)")
    tb9 = s9.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    tf9 = tb9.text_frame
    tf9.word_wrap = True

    sel_points = [
        ("Why Tabular ML Fails: ", "XGBoost and Random Forests cannot propagate high-order signals across disparate entities (e.g., weather impacting building envelope, causing appliance strain, triggering parts demand)."),
        ("Why Hetero-GNN Wins: ", "Relational Graph Convolution (RGCN / HeteroConv) learns relation-specific message transformation matrices W_r for each edge type."),
        ("Multi-Task Prediction Heads: ", "Latent representations feed 5 decoupled MLP heads: Demand (kW/ton), Power (kW), Failure (0/1), Sleeper Churn (0/1), and Inventory Velocity.")
    ]
    for b_title, b_desc in sel_points:
        p = tf9.add_paragraph()
        p.space_after = Pt(8)
        r1 = p.add_run()
        set_font(r1, b_title, "Times New Roman", 12, bold=True, color=(15, 34, 64))
        r2 = p.add_run()
        set_font(r2, b_desc, "Times New Roman", 11, color=(50, 50, 50))

    img_arch = "visualization/03_hetero_gnn_model_architecture.png"
    if os.path.exists(img_arch):
        s9.shapes.add_picture(img_arch, Inches(5.4), Inches(1.8), Inches(4.2), Inches(4.8))

    # -------------------------------------------------------------
    # SLIDE 10: Model Training — Tesla V100 GPU Setup (Animesh)
    # -------------------------------------------------------------
    s10 = add_content_slide("Chapter 8.1: Model Training on Tesla V100 GPU (Animesh)")
    tb10 = s10.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    tf10 = tb10.text_frame
    tf10.word_wrap = True

    train_points = [
        ("Enterprise GPU Infrastructure: ", "Trained on dedicated NVIDIA Tesla V100S-PCIE-32GB GPU server (10.13.1.7) with CUDA 13.0, PyTorch 2.6.0, and PyG 2.8.1."),
        ("Automatic Mixed Precision (AMP): ", "Utilized torch.cuda.amp.autocast and GradScaler for FP16 training, bounding peak VRAM allocation to only 28.8 MB."),
        ("Self-Healing Safeguards: ", "Automated divergence recovery catches NaN/Inf or loss spikes > 10x moving average, rolls back to best checkpoint, and decays learning rate by 50%."),
        ("120-Epoch Pre-Training: ", "Initial composite loss dropped rapidly from ~245.0 to a plateau of 137.04 (artifacts/best_hetero_gnn_v100.pt).")
    ]
    for b_title, b_desc in train_points:
        p = tf10.add_paragraph()
        p.space_after = Pt(8)
        r1 = p.add_run()
        set_font(r1, b_title, "Times New Roman", 11.5, bold=True, color=(15, 34, 64))
        r2 = p.add_run()
        set_font(r2, b_desc, "Times New Roman", 11, color=(50, 50, 50))

    img_heal = "visualization/05_self_healing_training_flowchart.png"
    if os.path.exists(img_heal):
        s10.shapes.add_picture(img_heal, Inches(5.4), Inches(1.8), Inches(4.2), Inches(4.8))

    # -------------------------------------------------------------
    # SLIDE 11: Model Evaluation & Audit Bottlenecks (Animesh)
    # -------------------------------------------------------------
    s11 = add_content_slide("Chapter 8.2: Model Evaluation & Audit Bottlenecks (Animesh)")
    tb11 = s11.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(5.0))
    tf11 = tb11.text_frame
    tf11.word_wrap = True

    audit_points = [
        ("The Argos Technical Audit Diagnosis: ", "Rigorous evaluation of the pre-trained checkpoint revealed 3 severe structural performance bottlenecks:"),
        ("1. Multi-Task Gradient Dominance: ", "Unweighted Demand MSE (~115) and Inventory MSE (~20) produced gradients 100x larger than Failure BCE (~0.70) and Churn BCE (~0.50). Shared GNN layers ignored critical component breakdowns."),
        ("2. Non-Peak Regime Under-Calibration: ", "During temperate and monsoon months (Feb, Mar, Jul-Nov), the model flagged exactly zero high-risk equipment units, completely failing to detect moisture-driven PCB corrosion."),
        ("3. Outlier Thermal Spikes: ", "Extreme 45°C summer heatwaves generated squared-error spikes (MSE) that destabilized gradient descent in later training epochs.")
    ]
    for b_title, b_desc in audit_points:
        p = tf11.add_paragraph()
        p.space_after = Pt(11)
        r1 = p.add_run()
        set_font(r1, b_title, "Times New Roman", 12.5, bold=True, color=(15, 34, 64))
        r2 = p.add_run()
        set_font(r2, b_desc, "Times New Roman", 11.5, color=(40, 40, 40))

    # -------------------------------------------------------------
    # SLIDE 12: Fine-Tuning: Kendall Uncertainty & Huber Loss (Animesh)
    # -------------------------------------------------------------
    s12 = add_content_slide("Chapter 8.2: Fine-Tuning Mathematical Innovations (Animesh)")
    tb12 = s12.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    tf12 = tb12.text_frame
    tf12.word_wrap = True

    ft_points = [
        ("Kendall Uncertainty Weighting: ", "Replaced manual grid search with learnable task log-variances s_k = log(sigma_k^2): L = sum_k [ 0.5*exp(-s_k)*L_k + 0.5*s_k ], automatically balancing multi-task gradients."),
        ("SmoothL1 (Huber) Loss: ", "Applied transition threshold beta=4.0 for regression heads. Replaces quadratic explosions with bounded constant gradients (|grad|=1.0) during summer heatwaves."),
        ("Real Benchmark Data Ingestion: ", "Ingested UCI AI4I 2020 (HDF/PWF faults), UCI Online Retail II (4,312 customer RFMs), and Kaggle HVAC (8,760h telemetry)."),
        ("V100 Execution: ", "60 epochs completed in 17.0s on Tesla V100 GPU (28.8MB VRAM).")
    ]
    for b_title, b_desc in ft_points:
        p = tf12.add_paragraph()
        p.space_after = Pt(7)
        r1 = p.add_run()
        set_font(r1, b_title, "Times New Roman", 11.5, bold=True, color=(15, 34, 64))
        r2 = p.add_run()
        set_font(r2, b_desc, "Times New Roman", 11, color=(50, 50, 50))

    img_loss = "visualization/06_pretrain_vs_finetune_loss_trajectory.png"
    if os.path.exists(img_loss):
        s12.shapes.add_picture(img_loss, Inches(5.4), Inches(2.0), Inches(4.2), Inches(4.4))

    # -------------------------------------------------------------
    # SLIDE 13: Before vs. After Calibration Ledger (Animesh)
    # -------------------------------------------------------------
    s13 = add_content_slide("Chapter 8.2: Quantitative Evaluation Ledger (Animesh)")
    tb13 = s13.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    tf13 = tb13.text_frame
    tf13.word_wrap = True

    ledger_bullets = [
        ("Overall Loss: ", "Plummeted from 137.04 to 46.80 (-65.85% total error reduction)."),
        ("Pillar 1 Demand MSE: ", "Reduced from 115.61 to 33.13 (-71.34%, MAE: 8.77 -> 4.64)."),
        ("Pillar 2 Power MSE: ", "Reduced from 0.8364 to 0.0654 (-92.18%, MAE: 0.749 -> 0.209 kW)."),
        ("Pillar 2 Failure F1: ", "Increased from 0.540 to 0.697 (+29.07%), proactively flagging 432 units across all months."),
        ("Pillar 3 Churn F1: ", "Jumped from 0.397 to 0.646 (+62.72%), identifying 278 sleepers and reactivating 276 customers."),
        ("Pillar 4 Inventory MSE: ", "Reduced from 19.13 to 12.61 (-34.11%, prevents stockouts).")
    ]
    for b_title, b_desc in ledger_bullets:
        p = tf13.add_paragraph()
        p.space_after = Pt(6)
        r1 = p.add_run()
        set_font(r1, b_title, "Times New Roman", 11, bold=True, color=(15, 34, 64))
        r2 = p.add_run()
        set_font(r2, b_desc, "Times New Roman", 10.5, color=(50, 50, 50))

    img_comp = "visualization/07_pillar_performance_comparison_before_after.png"
    if os.path.exists(img_comp):
        s13.shapes.add_picture(img_comp, Inches(5.4), Inches(2.0), Inches(4.2), Inches(4.4))

    # -------------------------------------------------------------
    # SLIDE 14: Model Deployment & Interactive Dashboard (Saurav)
    # -------------------------------------------------------------
    s14 = add_content_slide("Chapter 8.3: Model Deployment & Dashboard (Saurav)")
    tb14 = s14.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(5.0))
    tf14 = tb14.text_frame
    tf14.word_wrap = True

    dep_points = [
        ("Streamlit Cloud Production Deployment: ", "Deployed as an interactive executive command center accessible globally at https://smarthvac-advisor.streamlit.app/ with continuous cloud CI/CD."),
        ("In-Memory Zero-Crash Graph Rendering: ", "Engineered failsafe PyVis in-memory HTML injection (net.generate_html()), completely eliminating filesystem permission errors and blank screen crashes."),
        ("Clean 5-Tab Operational Interface: ", "Optimized dashboard architecture into 5 focused tabs: (1) Interactive Hetero-Graph, (2) Device Failure Telemetry, (3) Sleeper Retention Engine, (4) Spare Parts Forecasting, and (5) Executive Financial Ledger."),
        ("Standalone Offline Browser Dashboard: ", "Delivered ARGOS_BROWSER_RESULTS.html: a 100% self-contained, offline interactive HTML5 dashboard with zero cloud or server dependencies.")
    ]
    for b_title, b_desc in dep_points:
        p = tf14.add_paragraph()
        p.space_after = Pt(11)
        r1 = p.add_run()
        set_font(r1, b_title, "Times New Roman", 12.5, bold=True, color=(15, 34, 64))
        r2 = p.add_run()
        set_font(r2, b_desc, "Times New Roman", 11.5, color=(40, 40, 40))

    # -------------------------------------------------------------
    # SLIDE 15: Model Monitoring & Drift Maintenance (Saurav)
    # -------------------------------------------------------------
    s15 = add_content_slide("Chapter 8.4: Model Monitoring & Maintenance (Saurav)")
    tb15 = s15.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(5.0))
    tf15 = tb15.text_frame
    tf15.word_wrap = True

    mon_points = [
        ("Continuous Drift Monitoring: ", "Automated telemetry logger tracks real-time sensor streams and tests for two distinct forms of operational degradation:"),
        ("1. Covariate (Feature) Climate Drift: ", "Tracks statistical distribution shifts in ambient temperatures and humidity using Kolmogorov-Smirnov (KS) tests and Population Stability Index (PSI > 0.25)."),
        ("2. Concept Appliance Wear Drift: ", "Monitors degradation in physical relationships as aging equipment deviates from nominal factory power-COP curves over multi-year lifespans."),
        ("Automated Retraining Triggers: ", "If power prediction MSE exceeds 0.20 kW or sleeper churn F1 drops below 0.55, the pipeline automatically triggers incremental V100 fine-tuning (fine_tune_v100.py) over the latest seasonal window.")
    ]
    for b_title, b_desc in mon_points:
        p = tf15.add_paragraph()
        p.space_after = Pt(11)
        r1 = p.add_run()
        set_font(r1, b_title, "Times New Roman", 12.5, bold=True, color=(15, 34, 64))
        r2 = p.add_run()
        set_font(r2, b_desc, "Times New Roman", 11.5, color=(40, 40, 40))

    # -------------------------------------------------------------
    # SLIDE 16: Audited 12-Month Financial Ledger & Business Impact
    # -------------------------------------------------------------
    s16 = add_content_slide("Chapter 9: Audited 12-Month Financial Ledger")
    tb16 = s16.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    tf16 = tb16.text_frame
    tf16.word_wrap = True

    fin_bullets = [
        ("Gross Revenue: ", "₹2,61,01,851.84 (~₹2.61 Crore), gaining +₹20.78 Lakhs (+8.65%) via proactive sleeper reactivation."),
        ("Total OPEX: ", "₹34,00,492.40 (~₹34.00 Lakhs), funding 432 proactive component replacements."),
        ("Net Operating Profit: ", "₹2,27,01,359.44 (~₹2.27 Crore), maintaining 100% positive margin across all 12 operating months."),
        ("Cost-Efficiency Ratio: ", "7.764x mean / 7.676x aggregate annual ratio (₹7.68-₹7.76 returned per ₹1.00 spent)."),
        ("Proactive Breakdown Savings: ", "+22.64% average operational cost reduction (peaking at +28.4% in peak winter and summer).")
    ]
    for b_title, b_desc in fin_bullets:
        p = tf16.add_paragraph()
        p.space_after = Pt(7)
        r1 = p.add_run()
        set_font(r1, b_title, "Times New Roman", 11.5, bold=True, color=(15, 34, 64))
        r2 = p.add_run()
        set_font(r2, b_desc, "Times New Roman", 11, color=(50, 50, 50))

    img_fin = "visualization/08_12month_audited_financial_ledger.png"
    if os.path.exists(img_fin):
        s16.shapes.add_picture(img_fin, Inches(5.4), Inches(2.0), Inches(4.2), Inches(4.4))

    # -------------------------------------------------------------
    # SLIDE 17: Challenges, Future Work & Conclusion
    # -------------------------------------------------------------
    s17 = add_content_slide("Chapter 10 & 11: Challenges & Conclusion")
    tb17 = s17.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(5.0))
    tf17 = tb17.text_frame
    tf17.word_wrap = True

    concl_points = [
        ("Engineering Challenges Overcome: ", "Resolved multi-task gradient dominance via Kendall uncertainty loss, prevented thermal outlier explosions via Huber penalties, and eliminated class imbalance collapse via positive weighting."),
        ("Future Direction 1 — Continuous-Time Dynamic Graphs: ", "Extend static monthly graph snapshots into continuous-time dynamic graph neural networks (CTDG) with temporal attention (TGAT) to process sub-second sensor streams."),
        ("Future Direction 2 — Federated Privacy-Preserving GNN: ", "Deploy Federated GNN edge gateways to learn cross-facility degradation patterns while keeping tenant setpoints and billing histories strictly local and private."),
        ("Final Takeaway: ", "The Smart HVAC Hetero-GNN proves that unifying cyber-physical topology with business retention economics delivers superior predictive accuracy (-65.85% loss) and audited commercial profitability (7.764x cost efficiency).")
    ]
    for b_title, b_desc in concl_points:
        p = tf17.add_paragraph()
        p.space_after = Pt(11)
        r1 = p.add_run()
        set_font(r1, b_title, "Times New Roman", 12.5, bold=True, color=(15, 34, 64))
        r2 = p.add_run()
        set_font(r2, b_desc, "Times New Roman", 11.5, color=(40, 40, 40))

    # -------------------------------------------------------------
    # SLIDE 18: Thank You / Q&A Slide (Layout 4)
    # -------------------------------------------------------------
    s18 = prs.slides.add_slide(prs.slide_layouts[4])
    for shape in s18.shapes:
        if shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            set_font(r, "Thank You!\nQuestions & Discussion", "Times New Roman", 28, bold=True, color=(15, 34, 64))
            
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            r2 = p2.add_run()
            set_font(r2, "\nSmart HVAC Heterogeneous GML Project\nGitHub: https://github.com/AnimeshBasak-14/SmartHVAC\nLive Dashboard: https://smarthvac-advisor.streamlit.app/", "Times New Roman", 13, bold=False, color=(50, 50, 50))

    prs.save(out_path)
    print(f"Full Project PowerPoint generated successfully: {out_path} ({len(prs.slides)} slides)")

if __name__ == "__main__":
    build_presentation()
